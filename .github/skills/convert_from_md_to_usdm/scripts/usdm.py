#!/usr/bin/env python3
"""
USDM プロジェクト管理ツール

Usage:
    python3 usdm.py new <project_name> <output_dir>
    python3 usdm.py add <md_dir> <sheet_name>
"""

import argparse
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import Outline

# Import markdown parsers
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import md2usdm as m2u


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

DEFAULT_FONT = Font(name='ＭＳ ゴシック', size=9)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)
VCENTER = Alignment(vertical='center', wrap_text=True)
VCENTER_NOWRAP = Alignment(vertical='center')
HCENTER_VCENTER = Alignment(horizontal='center', vertical='center')

# Template fills (theme 9 tint 0.8 ≈ E2EFDA, theme 7 tint 0.8 ≈ FFF2CC)
U_FILL = PatternFill(start_color='FFE2EFDA', end_color='FFE2EFDA', fill_type='solid')
UU_FILL = PatternFill(start_color='FFFFF2CC', end_color='FFFFF2CC', fill_type='solid')

# Column widths from template
TEMPLATE_WIDTHS = {'A': 11.0, 'B': 5.25, 'C': 10.375, 'D': 20.625, 'E': 70.625}


# ---------------------------------------------------------------------------
# Markdown Generation
# ---------------------------------------------------------------------------

def gen_top_index(project_name, sheets):
    lines = [
        f"# {project_name}",
        "",
        "## ヘッダー",
        "",
        "| 行 | A | **B** | ~~C~~ | D | E |",
        "|---|---|---|---|---|---|",
        "| 列 | 11.0 | 5.25 | 10.375 | 20.625 | 70.625 |",
        "| 1 | カテゴリ名 | 要求 | 要求ID | 要求仕様 | ← |",
        "",
        "## シート一覧",
        "",
    ]
    for sheet_name, dir_name in sheets:
        lines.append(f"- [{sheet_name}](./{dir_name}/index.md)")
    lines.append("")
    return "\n".join(lines)


def gen_component_file(project_name):
    return "\n".join([
        f"# 【{project_name}】コンポーネント",
        "",
        "## 要求",
        "",
        "コンポーネント",
        "",
        "### 理由",
        "",
        "記載無し",
        "",
        "### 説明",
        "",
        "記載無し",
        "",
    ])


def gen_sheet_index(project_name, sheet_name, upper_reqs):
    """Generate sheet index.md.

    upper_reqs: [(req_id, title)] e.g. [("R01", "サンプル要件")]
    """
    lines = [
        f"↑ [{project_name}](../index.md)",
        "",
        f"# 【{project_name}】{sheet_name}",
        "",
        "## 要求",
        "",
        f"[{project_name}: コンポーネント](../{project_name}.md \"【{project_name}】\")",
        "",
        "## 下位要求",
        "",
    ]
    for req_id, title in upper_reqs:
        full_id = f"{project_name}_{req_id}"
        lines.append(
            f"- [{req_id} {title}](./{full_id}/{full_id}.md \"【{full_id}】\")"
        )
    lines.append("")
    return "\n".join(lines)


def gen_u_file(project_name, req_id, title, sheet_name, lower_reqs):
    """Generate u-file.

    lower_reqs: [(dm_id, title)] e.g. [("01", "サンプル下位要件")]
    """
    full_id = f"{project_name}_{req_id}"
    lines = [
        f"↑ [{sheet_name}](../index.md)",
        "",
        f"# 【{full_id}】{title}",
        "",
        "## 要求",
        "",
        "（要求内容を記述してください）",
        "",
        "### 理由",
        "",
        "（理由を記述してください）",
        "",
        "### 説明",
        "",
        "記載無し",
        "",
        "## 下位要求",
        "",
    ]
    for dm_id, dm_title in lower_reqs:
        dm_full = f"{full_id}.DM{dm_id}"
        lines.append(
            f"- [{dm_id} {dm_title}](./{dm_full}.md \"【{dm_full}】\")"
        )
    lines.append("")
    return "\n".join(lines)


def gen_uu_file(project_name, req_id, dm_id, dm_title, parent_title, parent_spec_id, specs):
    """Generate uu-file.

    specs: [(num, content)] e.g. [("01", "（仕様内容を記述してください）")]
    """
    full_id = f"{project_name}_{req_id}"
    dm_full = f"{full_id}.DM{dm_id}"
    lines = [
        f"↑ [【{full_id}】{parent_title}](./{full_id}.md \"{parent_spec_id}\")",
        "",
        f"# 【{dm_full}】{dm_title}",
        "",
        "## 要求",
        "",
        "（下位要求内容を記述してください）",
        "",
        "### 理由",
        "",
        "（理由を記述してください）",
        "",
        "### 説明",
        "",
        "記載無し",
        "",
        "## 仕様",
        "",
        "| 種別 | 内容 | 仕様番号 |",
        "| --- | --- | --- |",
    ]
    for num, content in specs:
        spec_id = f"【{dm_full}-{num}】"
        lines.append(f"| 仕様 | {content} | {spec_id} |")
    lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Excel Generation (Template Format)
# ---------------------------------------------------------------------------

def _set(ws, row, col, value, font=None, fill=None, alignment=None, border=None):
    """Set cell value and style."""
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = font or DEFAULT_FONT
    cell.alignment = alignment or VCENTER
    cell.border = border if border is not None else THIN_BORDER
    if fill:
        cell.fill = fill


def write_template_xlsx(md_dir, output_path=None):
    """Parse markdown and generate template-format Excel."""
    top_index = os.path.join(md_dir, "index.md")
    result = m2u.parse_top_index(top_index)
    header_entries, sheet_list, excel_basename = result[0], result[1], result[2]
    base_col_widths = result[3]

    if output_path is None:
        output_path = os.path.join(md_dir, f"{excel_basename}.xlsx")

    # Find component file
    comp = None
    for f in os.listdir(md_dir):
        if f.endswith(".md") and f != "index.md":
            comp = m2u.parse_component_file(os.path.join(md_dir, f))
            break

    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    for sheet_name, dir_name in sheet_list:
        # Parse sheet using m2u parsers (config values won't be used for writing)
        dummy_config = {
            "data_start_row": 2, "col_spec_id": None, "col_content": 5,
            "col_usdm_ref": None, "col_auto": None,
            "header_extra_cols": [], "col_widths": {},
            "extra_col_colors": ({}, {}),
            "has_sections": False, "section_marker": None,
        }
        sheet_data = _parse_template_sheet(md_dir, dir_name, sheet_name, comp)

        ws = wb.create_sheet(title=sheet_name)
        _write_template_sheet(ws, sheet_data, base_col_widths)

    wb.properties.title = excel_basename
    wb.save(output_path)
    return output_path


def _parse_template_sheet(md_dir, dir_name, sheet_name, comp):
    """Parse sheet from markdown, returning SheetData for template format."""
    dirpath = os.path.join(md_dir, dir_name)
    index_path = os.path.join(dirpath, "index.md")
    content = m2u.read_md(index_path)

    # Derive separator label
    category = sheet_name.split(".", 1)[1] if "." in sheet_name else sheet_name
    sep_label = f"＜{category}＞"

    comp_reason = comp.reason if comp else ""
    comp_explanation = comp.explanation if comp else ""
    comp_info = comp.info if comp else ""

    # Minimal config for parsing u/uu files
    config = {
        "data_start_row": 2, "col_spec_id": None, "col_content": 5,
        "col_usdm_ref": None, "col_auto": None,
        "header_extra_cols": [], "col_widths": {},
        "extra_col_colors": ({}, {}),
        "has_sections": False, "section_marker": None,
    }

    sheet = m2u.SheetData(
        sheet_name=sheet_name, config=config,
        component_info=comp_info,
        component_reason=comp_reason,
        component_explanation=comp_explanation,
        separator_label=sep_label,
    )

    # Detect sections vs flat
    lines = content.split("\n")
    sections_found = []
    current_section_name = None
    current_links = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("## ") and stripped not in (
            "## ヘッダー", "## 上位要件", "## 下位要件", "## 上位要求", "## 下位要求", "## シート一覧", "## 要件", "## 要求"
        ):
            if current_section_name is not None:
                sections_found.append((current_section_name, current_links))
            current_section_name = stripped[3:].strip()
            current_links = []
            continue
        if stripped in ("## 上位要件", "## 下位要件", "## 上位要求", "## 下位要求"):
            current_section_name = None
            current_links = []
            continue
        link_m = re.match(r'- \[(.+?)\]\((.+?)\.md(?:\s+"(.+?)")?\)', stripped)
        if link_m:
            current_links.append((link_m.group(1), link_m.group(2) + ".md", link_m.group(3) or ""))
    if current_section_name is not None:
        sections_found.append((current_section_name, current_links))

    if sections_found:
        config["has_sections"] = True
        config["section_marker"] = "t"
        if re.search(r'セクション区分: `tt`', content):
            config["section_marker"] = "tt"
        for sec_name, links in sections_found:
            num = sec_name.split(".")[0] if "." in sec_name else sec_name
            section = m2u.Section(number=num, name=sec_name)
            for _, link_path, _ in links:
                full_path = os.path.join(dirpath, link_path)
                if os.path.exists(full_path):
                    u = m2u.parse_u_file(full_path, os.path.dirname(full_path), config)
                    section.upper_reqs.append(u)
            sheet.sections.append(section)
    else:
        for line in lines:
            stripped = line.strip()
            link_m = re.match(r'- \[(.+?)\]\((.+?)\.md(?:\s+"(.+?)")?\)', stripped)
            if link_m:
                link_path = link_m.group(2) + ".md"
                full_path = os.path.join(dirpath, link_path)
                if os.path.exists(full_path):
                    u = m2u.parse_u_file(full_path, os.path.dirname(full_path), config)
                    sheet.upper_reqs.append(u)

    return sheet


def _write_template_sheet(ws, sheet_data, col_widths):
    """Write a template-format sheet."""
    # Set column widths
    for letter, width in (col_widths or TEMPLATE_WIDTHS).items():
        ws.column_dimensions[letter].width = width

    # Row 1: Header
    _set(ws, 1, 1, "カテゴリ名", alignment=HCENTER_VCENTER)
    _set(ws, 1, 2, "要求", alignment=HCENTER_VCENTER)
    _set(ws, 1, 3, "要求ID", alignment=HCENTER_VCENTER)
    _set(ws, 1, 4, "要求仕様", alignment=HCENTER_VCENTER)
    _set(ws, 1, 5, "", alignment=HCENTER_VCENTER)
    ws.merge_cells('D1:E1')

    # Outline
    ws.sheet_properties.outlinePr = Outline(summaryBelow=True, summaryRight=True)

    # Data rows
    row = 2
    has_sections = sheet_data.config["has_sections"]

    if has_sections:
        for section in sheet_data.sections:
            # Section header in column A
            for u in section.upper_reqs:
                row = _write_template_u(ws, row, u, section.name)
    else:
        for u in sheet_data.upper_reqs:
            row = _write_template_u(ws, row, u, "")

    # Auto-filter
    ws.auto_filter.ref = f"A1:E{max(row - 1, 1)}"

    return row


def _write_template_u(ws, row, u, section_name):
    """Write u + o + o rows (upper requirement)."""
    value, font = m2u.md_to_cell(u.text)

    # u row: B=要求, C=req_id, D-E=content
    _set(ws, row, 1, section_name, fill=U_FILL)
    _set(ws, row, 2, "要求", fill=U_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 3, u.req_id, font=font, fill=U_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 4, value, font=font, fill=U_FILL)
    _set(ws, row, 5, "", fill=U_FILL)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
    ws.row_dimensions[row].outline_level = 0
    row += 1

    # o row: C=理由, D-E=reason
    reason = u.reason or "記載無し"
    r_value, r_font = m2u.md_to_cell(reason)
    _set(ws, row, 1, "", fill=U_FILL)
    _set(ws, row, 2, "", fill=U_FILL)
    _set(ws, row, 3, "理由", fill=U_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 4, r_value, font=r_font, fill=U_FILL)
    _set(ws, row, 5, "", fill=U_FILL)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
    ws.row_dimensions[row].outline_level = 1
    row += 1

    # o row: C=説明, D-E=explanation
    expl = u.explanation or "記載無し"
    e_value, e_font = m2u.md_to_cell(expl)
    _set(ws, row, 1, "", fill=U_FILL)
    _set(ws, row, 2, "", fill=U_FILL)
    _set(ws, row, 3, "説明", fill=U_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 4, e_value, font=e_font, fill=U_FILL)
    _set(ws, row, 5, "", fill=U_FILL)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
    ws.row_dimensions[row].outline_level = 1
    row += 1

    # Lower requirements
    for uu in u.lower_reqs:
        row = _write_template_uu(ws, row, uu)

    return row


def _write_template_uu(ws, row, uu):
    """Write uu + oo + oo + ss rows (lower requirement)."""
    value, font = m2u.md_to_cell(uu.text)

    # uu row: C=要求, D=req_id, E=content
    _set(ws, row, 1, "")
    _set(ws, row, 2, "")
    _set(ws, row, 3, "要求", fill=UU_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 4, uu.req_id, font=font, fill=UU_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 5, value, font=font, fill=UU_FILL)
    ws.row_dimensions[row].outline_level = 1
    row += 1

    # oo row: D=理由, E=reason
    reason = uu.reason or "記載無し"
    r_value, r_font = m2u.md_to_cell(reason)
    _set(ws, row, 1, "")
    _set(ws, row, 2, "")
    _set(ws, row, 3, "", fill=UU_FILL)
    _set(ws, row, 4, "理由", fill=UU_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 5, r_value, font=r_font, fill=UU_FILL)
    ws.row_dimensions[row].outline_level = 2
    row += 1

    # oo row: D=説明, E=explanation
    expl = uu.explanation or "記載無し"
    e_value, e_font = m2u.md_to_cell(expl)
    _set(ws, row, 1, "")
    _set(ws, row, 2, "")
    _set(ws, row, 3, "", fill=UU_FILL)
    _set(ws, row, 4, "説明", fill=UU_FILL, alignment=HCENTER_VCENTER)
    _set(ws, row, 5, e_value, font=e_font, fill=UU_FILL)
    ws.row_dimensions[row].outline_level = 2
    row += 1

    # Specifications
    for spec in uu.specs:
        row = _write_template_ss(ws, row, spec)

    # tt subsections
    for tt in uu.subsections:
        _set(ws, row, 1, "")
        _set(ws, row, 2, "")
        _set(ws, row, 3, "")
        _set(ws, row, 4, tt.name)
        _set(ws, row, 5, "")
        ws.row_dimensions[row].outline_level = 2
        row += 1
        for spec in tt.specs:
            row = _write_template_ss(ws, row, spec)

    return row


def _write_template_ss(ws, row, spec):
    """Write ss row (specification)."""
    value, font = m2u.md_to_cell(spec.content)

    _set(ws, row, 1, "")
    _set(ws, row, 2, "")
    _set(ws, row, 3, spec.ss_type, alignment=HCENTER_VCENTER)
    _set(ws, row, 4, spec.spec_id)
    _set(ws, row, 5, value, font=font)
    ws.row_dimensions[row].outline_level = 2
    return row + 1


# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------

def cmd_new(project_name, output_dir):
    """新規 USDM プロジェクト作成"""
    os.makedirs(output_dir, exist_ok=True)

    # Default: 1 sheet, 1 upper req, 1 lower req, 1 spec
    sheets = [("1.要件", "sheet01")]
    req_id = "R01"
    req_title = "サンプル要件"
    dm_id = "01"
    dm_title = "サンプル下位要件"

    # index.md
    index_content = gen_top_index(project_name, sheets)
    _write_file(os.path.join(output_dir, "index.md"), index_content)

    # component file
    comp_content = gen_component_file(project_name)
    _write_file(os.path.join(output_dir, f"{project_name}.md"), comp_content)

    # sheet01/
    sheet_dir = os.path.join(output_dir, "sheet01")
    os.makedirs(sheet_dir, exist_ok=True)

    sheet_index = gen_sheet_index(project_name, "1.要件", [(req_id, req_title)])
    _write_file(os.path.join(sheet_dir, "index.md"), sheet_index)

    # u-file directory
    full_id = f"{project_name}_{req_id}"
    u_dir = os.path.join(sheet_dir, full_id)
    os.makedirs(u_dir, exist_ok=True)

    u_content = gen_u_file(project_name, req_id, req_title, "1.要件",
                           [(dm_id, dm_title)])
    _write_file(os.path.join(u_dir, f"{full_id}.md"), u_content)

    uu_content = gen_uu_file(project_name, req_id, dm_id, dm_title,
                             req_title, f"【{full_id}】",
                             [("01", "（仕様内容を記述してください）")])
    _write_file(os.path.join(u_dir, f"{full_id}.DM{dm_id}.md"), uu_content)

    # Generate xlsx
    xlsx_path = os.path.join(output_dir, f"{project_name}.xlsx")
    write_template_xlsx(output_dir, xlsx_path)

    print(f"プロジェクト作成完了: {output_dir}/")
    print(f"  index.md")
    print(f"  {project_name}.md")
    print(f"  {project_name}.xlsx")
    print(f"  sheet01/index.md")
    print(f"  sheet01/{full_id}/{full_id}.md")
    print(f"  sheet01/{full_id}/{full_id}.DM{dm_id}.md")


def cmd_add(md_dir, sheet_name):
    """既存プロジェクトにシート追加"""
    index_path = os.path.join(md_dir, "index.md")
    if not os.path.exists(index_path):
        print(f"ERROR: {index_path} が見つかりません")
        sys.exit(1)

    content = m2u.read_md(index_path)

    # Get project name from title
    m = re.match(r'# (.+)', content)
    project_name = m.group(1).strip() if m else "Project"

    # Get existing sheet list
    existing = re.findall(r'- \[.+?\]\(\./(.+?)/index\.md\)', content)
    next_num = len(existing) + 1
    dir_name = f"sheet{next_num:02d}"

    # Create sheet directory
    sheet_dir = os.path.join(md_dir, dir_name)
    os.makedirs(sheet_dir, exist_ok=True)

    # Default: 1 upper req
    req_id = "R01"
    req_title = "サンプル要件"
    dm_id = "01"
    dm_title = "サンプル下位要件"

    sheet_index = gen_sheet_index(project_name, sheet_name, [(req_id, req_title)])
    _write_file(os.path.join(sheet_dir, "index.md"), sheet_index)

    full_id = f"{project_name}_{req_id}"
    u_dir = os.path.join(sheet_dir, full_id)
    os.makedirs(u_dir, exist_ok=True)

    u_content = gen_u_file(project_name, req_id, req_title, sheet_name,
                           [(dm_id, dm_title)])
    _write_file(os.path.join(u_dir, f"{full_id}.md"), u_content)

    uu_content = gen_uu_file(project_name, req_id, dm_id, dm_title,
                             req_title, f"【{full_id}】",
                             [("01", "（仕様内容を記述してください）")])
    _write_file(os.path.join(u_dir, f"{full_id}.DM{dm_id}.md"), uu_content)

    # Update top index.md: add new sheet to list
    new_line = f"- [{sheet_name}](./{dir_name}/index.md)"
    content = content.rstrip("\n")
    content += "\n" + new_line + "\n"
    _write_file(index_path, content)

    # Regenerate xlsx
    xlsx_path = os.path.join(md_dir, f"{project_name}.xlsx")
    write_template_xlsx(md_dir, xlsx_path)

    print(f"シート追加完了: {dir_name}/ ({sheet_name})")
    print(f"  {dir_name}/index.md")
    print(f"  {dir_name}/{full_id}/{full_id}.md")
    print(f"  {dir_name}/{full_id}/{full_id}.DM{dm_id}.md")
    print(f"  xlsx 再生成: {xlsx_path}")


def _write_file(path, content):
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="USDM プロジェクト管理ツール",
    )
    subparsers = parser.add_subparsers(dest="command")

    # new
    p_new = subparsers.add_parser("new", help="新規プロジェクト作成")
    p_new.add_argument("project_name", help="プロジェクト名 (例: PROJECT_ID)")
    p_new.add_argument("output_dir", help="出力ディレクトリ")

    # add
    p_add = subparsers.add_parser("add", help="シート追加")
    p_add.add_argument("md_dir", help="Markdown ディレクトリ")
    p_add.add_argument("sheet_name", help="シート名 (例: 2.機能要件)")

    args = parser.parse_args()

    if args.command == "new":
        cmd_new(args.project_name, args.output_dir)
    elif args.command == "add":
        cmd_add(args.md_dir, args.sheet_name)
    else:
        parser.print_help()
        sys.exit(1)


if __name__ == "__main__":
    main()
