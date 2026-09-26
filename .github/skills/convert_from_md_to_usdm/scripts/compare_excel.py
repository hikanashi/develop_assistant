#!/usr/bin/env python3
"""Compare original Excel with regenerated Excel cell by cell.

Usage:
    python3 compare_excel.py <original.xlsx> <regenerated.xlsx>
"""
import sys
import openpyxl
from openpyxl.cell.rich_text import CellRichText


DATA_START_ROW = 9


def is_usdm_sheet(ws):
    """Check if a worksheet is a USDM sheet by looking for markers in column B."""
    markers = {"t", "tt", "u", "o", "uu", "oo", "ss"}
    for row in range(DATA_START_ROW, min(DATA_START_ROW + 50, ws.max_row + 1)):
        v = ws.cell(row=row, column=2).value
        if v and str(v).strip() in markers:
            return True
    return False


def cell_value_str(cell):
    """Convert cell value to comparable string."""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        return str(v)
    return str(v)


def cell_font_summary(cell):
    """Summarize font properties."""
    f = cell.font
    if f is None:
        return ""
    parts = []
    if f.name:
        parts.append(f.name)
    if f.size:
        parts.append(f"sz={f.size}")
    if f.bold:
        parts.append("bold")
    if f.color and f.color.rgb and f.color.rgb != '00000000':
        parts.append(f"color={f.color.rgb}")
    if f.strike:
        parts.append("strike")
    return ",".join(parts)


def cell_fill_color(cell):
    """Get fill color."""
    f = cell.fill
    if f is None or f.fill_type is None:
        return ""
    if f.start_color and f.start_color.rgb:
        return f.start_color.rgb
    return ""


def cell_border_summary(cell):
    """Summarize border."""
    b = cell.border
    if b is None:
        return ""
    sides = []
    for name in ('left', 'right', 'top', 'bottom'):
        s = getattr(b, name)
        if s and s.style:
            sides.append(f"{name[0]}={s.style}")
    return ",".join(sides)


def is_merged(ws, row, col):
    """Check if a cell is part of a merged range (not the top-left)."""
    for mr in ws.merged_cells.ranges:
        if (mr.min_row <= row <= mr.max_row and
            mr.min_col <= col <= mr.max_col):
            if row == mr.min_row and col == mr.min_col:
                return False, str(mr)  # top-left of merge
            return True, str(mr)  # merged away
    return False, None


def compare_sheets(ws_orig, ws_regen, sheet_name):
    """Compare two worksheets and report differences."""
    diffs = []

    max_row = max(ws_orig.max_row, ws_regen.max_row)
    max_col = max(ws_orig.max_column, ws_regen.max_column)

    # Compare all rows (including header rows 1-8, since we now generate them)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            merged_orig, mr_orig = is_merged(ws_orig, row, col)
            merged_regen, mr_regen = is_merged(ws_regen, row, col)

            # Skip merged-away cells
            if merged_orig and merged_regen:
                continue

            c_orig = ws_orig.cell(row=row, column=col)
            c_regen = ws_regen.cell(row=row, column=col)

            v_orig = cell_value_str(c_orig)
            v_regen = cell_value_str(c_regen)

            # Collect all differences for this cell
            cell_diffs = []

            if v_orig != v_regen:
                cell_diffs.append(f"value: [{v_orig!r}] → [{v_regen!r}]")

            fill_orig = cell_fill_color(c_orig)
            fill_regen = cell_fill_color(c_regen)
            if fill_orig != fill_regen:
                cell_diffs.append(f"fill: [{fill_orig}] → [{fill_regen}]")

            border_orig = cell_border_summary(c_orig)
            border_regen = cell_border_summary(c_regen)
            if border_orig != border_regen:
                cell_diffs.append(f"border: [{border_orig}] → [{border_regen}]")

            if merged_orig != merged_regen:
                cell_diffs.append(f"merge: orig={mr_orig} regen={mr_regen}")

            if cell_diffs:
                col_letter = openpyxl.utils.get_column_letter(col)
                marker = ""
                if row >= DATA_START_ROW:
                    marker = cell_value_str(ws_orig.cell(row=row, column=2)) or "?"
                else:
                    marker = f"hdr"
                for d in cell_diffs:
                    diffs.append((row, col_letter, marker, d))

    return diffs


def main():
    if len(sys.argv) < 3:
        print("Usage: python3 compare_excel.py <original.xlsx> <regenerated.xlsx>")
        sys.exit(1)

    orig_path = sys.argv[1]
    regen_path = sys.argv[2]

    print(f"Loading original: {orig_path}")
    wb_orig = openpyxl.load_workbook(orig_path, rich_text=True)
    print(f"Loading regenerated: {regen_path}")
    wb_regen = openpyxl.load_workbook(regen_path, rich_text=True)

    # Find matching USDM sheets
    orig_sheets = [name for name in wb_orig.sheetnames if is_usdm_sheet(wb_orig[name])]
    regen_sheets = [name for name in wb_regen.sheetnames if is_usdm_sheet(wb_regen[name])]

    print(f"Original USDM sheets: {orig_sheets}")
    print(f"Regenerated USDM sheets: {regen_sheets}")

    total_diffs = 0
    for sheet_name in orig_sheets:
        if sheet_name not in wb_regen.sheetnames:
            print(f"\n[{sheet_name}] MISSING in regenerated file!")
            continue

        ws_orig = wb_orig[sheet_name]
        ws_regen = wb_regen[sheet_name]

        diffs = compare_sheets(ws_orig, ws_regen, sheet_name)

        if diffs:
            print(f"\n{'='*70}")
            print(f"[{sheet_name}] {len(diffs)} differences")
            print(f"{'='*70}")

            # Group by row
            from itertools import groupby
            for row_num, row_diffs in groupby(diffs, key=lambda x: x[0]):
                row_diffs = list(row_diffs)
                marker = row_diffs[0][2]
                print(f"  Row {row_num} ({marker}):")
                for _, col_letter, _, desc in row_diffs:
                    print(f"    {col_letter}: {desc}")
            total_diffs += len(diffs)
        else:
            print(f"[{sheet_name}] OK - no differences")

    print(f"\n{'='*70}")
    print(f"Total: {total_diffs} differences")
    print(f"{'='*70}")


if __name__ == "__main__":
    main()
